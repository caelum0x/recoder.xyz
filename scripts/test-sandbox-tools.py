"""Test sandbox tool execution — verifies run_code and create_file auto-convert work."""
import asyncio
import base64
import hashlib
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

from daytona_sdk import Daytona, DaytonaConfig


def test():
    config = DaytonaConfig()
    d = Daytona(config)
    sandboxes = d.list()
    items = sandboxes.items

    sb = None
    for s in items:
        state = str(getattr(s, 'state', ''))
        if 'STARTED' in state:
            sb = s
            break

    if not sb:
        print('FAIL: No running sandbox found')
        return False

    print(f'Sandbox: {sb.id} (state: {sb.state})')
    all_pass = True

    # Test 1: Basic exec
    print('\n=== Test 1: Basic exec ===')
    r = sb.process.exec('echo "hello world"', timeout=10)
    print(f'  exit_code={r.exit_code}, output={repr(r.result[:100])}')
    if r.exit_code != 0 or 'hello' not in (r.result or ''):
        print('  FAIL: basic exec')
        all_pass = False
    else:
        print('  PASS')

    # Test 2: Python available
    print('\n=== Test 2: Python available ===')
    r = sb.process.exec('python3 --version', timeout=10)
    print(f'  exit_code={r.exit_code}, output={repr(r.result[:100])}')
    if r.exit_code != 0:
        print('  FAIL: python3 not available')
        all_pass = False
    else:
        print('  PASS')

    # Test 3: pip install python-docx
    print('\n=== Test 3: pip install python-docx ===')
    r = sb.process.exec('pip install -q python-docx 2>&1 | tail -3', timeout=60)
    print(f'  exit_code={r.exit_code}, output={repr(r.result[:200])}')
    if r.exit_code != 0:
        print('  FAIL: pip install')
        all_pass = False
    else:
        print('  PASS')

    # Test 4: run_code path (write temp file + execute)
    print('\n=== Test 4: run_code path ===')
    code = 'print("Hello from run_code!")\nimport sys\nprint(f"Python {sys.version}")'
    code_hash = hashlib.md5(code.encode()).hexdigest()[:8]
    tmp_file = f'/tmp/_code_{code_hash}.py'

    # Try upload_file first
    try:
        sb.fs.upload_file(tmp_file, code.encode('utf-8'))
        print(f'  upload_file: OK')
    except Exception as e:
        print(f'  upload_file FAILED: {e}')
        b64 = base64.b64encode(code.encode('utf-8')).decode('ascii')
        sb.process.exec(f"echo '{b64}' | base64 -d > {tmp_file}", timeout=10)
        print(f'  fallback base64 write: OK')

    r = sb.process.exec(f'python3 {tmp_file}', timeout=30)
    print(f'  exit_code={r.exit_code}, output={repr(r.result[:200])}')
    sb.process.exec(f'rm -f {tmp_file}', timeout=5)
    if r.exit_code != 0 or 'Hello from run_code' not in (r.result or ''):
        print('  FAIL: run_code execution')
        all_pass = False
    else:
        print('  PASS')

    # Test 5: DOCX creation (auto-convert path)
    print('\n=== Test 5: DOCX auto-convert ===')
    text_content = '# Planets Report\nThe solar system has 8 planets.\n## Mercury\nClosest to the sun.\n## Jupiter\nThe largest planet.'
    full_path = '/workspace/test_planets.docx'

    b64_content = base64.b64encode(text_content.encode('utf-8')).decode('ascii')
    sb.process.exec(f"echo '{b64_content}' | base64 -d > /tmp/_docgen_content.txt", timeout=10)

    script = f"""
from docx import Document
doc = Document()
with open('/tmp/_docgen_content.txt', 'r') as f:
    content = f.read()
for para in content.split('\\n'):
    p = para.strip()
    if not p:
        continue
    if p.startswith('# '):
        doc.add_heading(p[2:], level=1)
    elif p.startswith('## '):
        doc.add_heading(p[3:], level=2)
    else:
        doc.add_paragraph(p)
doc.save('{full_path}')
import os
print(f'OK size={{os.path.getsize("{full_path}")}}')
"""
    script_b64 = base64.b64encode(script.encode('utf-8')).decode('ascii')
    sb.process.exec(f"echo '{script_b64}' | base64 -d > /tmp/_docgen_script.py", timeout=10)

    r = sb.process.exec('python3 /tmp/_docgen_script.py', timeout=30)
    print(f'  exit_code={r.exit_code}, output={repr(r.result[:200])}')
    if r.exit_code != 0 or 'OK' not in (r.result or ''):
        print('  FAIL: DOCX creation')
        all_pass = False
    else:
        print('  PASS')

    # Verify file
    r = sb.process.exec(f'test -s {full_path} && echo EXISTS || echo MISSING', timeout=10)
    print(f'  File check: {r.result.strip()}')

    # Test 6: ToolResult simulation
    print('\n=== Test 6: ToolResult success field ===')
    from dataclasses import dataclass
    from typing import Any

    @dataclass
    class ToolResult:
        success: bool
        output: Any

    result = ToolResult(success=True, output='Created test_planets.docx (real DOCX binary format)')
    print(f'  success type: {type(result.success)}, value: {result.success}')
    print(f'  success === True: {result.success is True}')

    import json
    serialized = json.dumps({"success": result.success, "output": result.output})
    deserialized = json.loads(serialized)
    print(f'  After JSON roundtrip: type={type(deserialized["success"])}, value={deserialized["success"]}')
    if deserialized["success"] is not True:
        print('  FAIL: JSON roundtrip corrupts boolean')
        all_pass = False
    else:
        print('  PASS: JSON preserves boolean correctly')

    print(f'\n{"=" * 50}')
    print(f'RESULT: {"ALL TESTS PASSED" if all_pass else "SOME TESTS FAILED"}')
    return all_pass


test()
